"""
Generador de documentos .docx de seguimiento bibliográfico.
Produce un archivo por cada artículo con estructura fija y estilos mínimos.
"""

import logging
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from config.settings import (
    DATA_OUTPUT_DIR,
    DOCX_FONT_NAME,
    DOCX_FONT_SIZE_HEADING1,
    DOCX_FONT_SIZE_HEADING2,
    DOCX_FONT_SIZE_NORMAL,
    PLACEHOLDER_TEXT,
)

logger = logging.getLogger("bibliografias")

_NOT_AVAILABLE = "No disponible"


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    """Agrega un encabezado con formato personalizado."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = DOCX_FONT_NAME
        if level == 1:
            run.font.size = Pt(DOCX_FONT_SIZE_HEADING1)
        else:
            run.font.size = Pt(DOCX_FONT_SIZE_HEADING2)


def _add_field(doc: Document, label: str, value: str) -> None:
    """Agrega un campo con etiqueta en negrita y valor normal."""
    para = doc.add_paragraph()
    run_label = para.add_run(f"{label}: ")
    run_label.bold = True
    run_label.font.name = DOCX_FONT_NAME
    run_label.font.size = Pt(DOCX_FONT_SIZE_NORMAL)
    run_value = para.add_run(value or _NOT_AVAILABLE)
    run_value.font.name = DOCX_FONT_NAME
    run_value.font.size = Pt(DOCX_FONT_SIZE_NORMAL)


def _add_paragraph(doc: Document, text: str) -> None:
    """Agrega un párrafo con formato estándar."""
    para = doc.add_paragraph(text or _NOT_AVAILABLE)
    for run in para.runs:
        run.font.name = DOCX_FONT_NAME
        run.font.size = Pt(DOCX_FONT_SIZE_NORMAL)


def _format_authors_ieee(authors: list[str]) -> str:
    """Formatea lista de autores al estilo IEEE: 'A. Apellido, B. Apellido, ...'"""
    ieee_authors = []
    for author in authors:
        parts = author.split(",")
        if len(parts) == 2:
            lastname = parts[0].strip()
            firstname = parts[1].strip()
            initials = ". ".join(
                n[0].upper() for n in firstname.split() if n
            )
            ieee_authors.append(f"{initials}. {lastname}")
        else:
            ieee_authors.append(author)
    if len(ieee_authors) > 3:
        return ", ".join(ieee_authors[:3]) + " et al."
    return ", ".join(ieee_authors)


def _build_ieee_reference(metadata: dict[str, Any]) -> str:
    """Construye la referencia bibliográfica en formato IEEE."""
    authors = _format_authors_ieee(metadata.get("autores", []))
    title = metadata.get("titulo", "Sin título")
    journal = metadata.get("journal", "")
    year = metadata.get("anio", "s.f.")
    doi = metadata.get("doi", "")

    ref = f'{authors}, "{title}"'
    if journal:
        ref += f", {journal}"
    ref += f", {year}."
    if doi:
        ref += f" DOI: {doi}"
    return ref


def generate_docx(
    metadata: dict[str, Any],
    analysis: dict[str, str] | None,
    codigo: str,
    output_dir: str | Path | None = None,
) -> Path:
    """
    Genera un documento .docx con la estructura de seguimiento bibliográfico.
    Si analysis es None (modo sin_llm), usa placeholders.
    Retorna la ruta del archivo generado.
    """
    output_dir = Path(output_dir) if output_dir else DATA_OUTPUT_DIR
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"{codigo}.docx"

    doc = Document()

    # Título principal
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run(f"Ficha de Seguimiento Bibliográfico — {codigo}")
    run.bold = True
    run.font.name = DOCX_FONT_NAME
    run.font.size = Pt(DOCX_FONT_SIZE_HEADING1 + 2)

    doc.add_paragraph()  # Separador

    # --- 1. Datos Generales ---
    _add_heading(doc, "1. Datos Generales", level=1)
    _add_field(doc, "Código", codigo)
    _add_field(doc, "Base de datos", metadata.get("base_datos", ""))
    _add_field(doc, "DOI", metadata.get("doi", ""))
    _add_field(doc, "URL", metadata.get("url", ""))
    _add_field(doc, "Fuente de metadatos", metadata.get("fuente_metadata", ""))

    # --- 2. Referencia Bibliográfica (IEEE) ---
    _add_heading(doc, "2. Referencia Bibliográfica (Formato IEEE)", level=1)
    ieee_ref = _build_ieee_reference(metadata)
    _add_paragraph(doc, ieee_ref)

    # --- 3. Información del Documento ---
    _add_heading(doc, "3. Información del Documento", level=1)
    _add_field(doc, "Título", metadata.get("titulo", ""))
    authors_str = "; ".join(metadata.get("autores", []))
    _add_field(doc, "Autores", authors_str)
    _add_field(doc, "Año", metadata.get("anio", ""))
    _add_field(doc, "Revista / Conferencia", metadata.get("journal", ""))
    _add_field(doc, "Tipo de documento", metadata.get("tipo", ""))

    _add_heading(doc, "Abstract", level=2)
    abstract = metadata.get("abstract", "")
    _add_paragraph(doc, abstract if abstract else _NOT_AVAILABLE)

    # --- 4. Análisis del Contenido ---
    _add_heading(doc, "4. Análisis del Contenido", level=1)

    use_llm = analysis is not None
    section_map = {
        "PROBLEMA": "Problema que aborda",
        "METODOLOGIA": "Metodología utilizada",
        "RESULTADOS": "Resultados principales",
        "APORTES": "Aportes relevantes",
        "LIMITACIONES": "Limitaciones identificadas",
    }
    for key, label in section_map.items():
        _add_heading(doc, label, level=2)
        text = analysis.get(key, "") if use_llm else ""
        _add_paragraph(doc, text if text else PLACEHOLDER_TEXT)

    # --- 5. Relación con mi Proyecto ---
    _add_heading(doc, "5. Relación con mi Proyecto", level=1)
    rel_text = analysis.get("RELACION_PROYECTO", "") if use_llm else ""
    _add_paragraph(doc, rel_text if rel_text else PLACEHOLDER_TEXT)

    # --- 6. Clasificación del Artículo ---
    _add_heading(doc, "6. Clasificación del Artículo", level=1)
    cls_text = analysis.get("CLASIFICACION", "") if use_llm else ""
    _add_paragraph(doc, cls_text if cls_text else PLACEHOLDER_TEXT)

    # Nota de error si el LLM falló
    if use_llm and analysis.get("_error"):
        doc.add_paragraph()
        error_para = doc.add_paragraph()
        run = error_para.add_run(f"Nota: Error en análisis LLM — {analysis['_error']}")
        run.italic = True
        run.font.name = DOCX_FONT_NAME
        run.font.size = Pt(9)

    doc.save(str(output_path))
    logger.info("Documento generado: %s", output_path.name)
    return output_path
